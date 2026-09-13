"""Per-doctor revenue/billing summary ("zestawienie"), exported as .xlsx.

Input is the dedicated zestawienie report template (spanning a week or
month), which - unlike the reminders template - carries a per-visit Status
column. Status decides whether a visit counts toward the doctor's total:
doctors/psychiatrists only count "Wykonane"; everyone else counts "Wykonane"
or "Nie zrealizowane". "Rezygnacja z wykonania" never counts. Visits
excluded by this rule are dropped from the export entirely, not just zeroed.
"""

import re
from pathlib import Path

import openpyxl

import employees_store
import report_parser

_ILLEGAL_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")

_DOCTOR_COUNTED_STATUSES = {"wykonane"}
_OTHER_COUNTED_STATUSES = {"wykonane", "nie zrealizowane"}


def _safe_sheet_name(name: str, used: set[str]) -> str:
    cleaned = _ILLEGAL_SHEET_CHARS.sub("", name).strip() or "Nieznany"
    cleaned = cleaned[:31]
    candidate = cleaned
    n = 2
    while candidate in used:
        suffix = f" ({n})"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate)
    return candidate


def _status_counts(status: str, is_psychiatra: bool) -> bool:
    allowed = _DOCTOR_COUNTED_STATUSES if is_psychiatra else _OTHER_COUNTED_STATUSES
    return status.strip().lower() in allowed


def generate(input_path: Path | str, output_path: Path | str, log=lambda msg: None) -> None:
    visits = report_parser.load_report(input_path)

    if not any(v.status for v in visits):
        raise ValueError(
            "Wybrany plik nie zawiera kolumny Status - wybierz raport w formacie "
            "przeznaczonym do zestawien."
        )

    employees = employees_store.load()
    warned_doctors: set[str] = set()

    by_doctor: dict[str, list] = {}
    for visit in visits:
        employee = employees_store.find_employee(visit.doctor, employees)
        if employee is None:
            if visit.doctor not in warned_doctors:
                warned_doctors.add(visit.doctor)
                log(
                    f"Nierozpoznany lekarz '{visit.doctor}' - zastosowano domyslna "
                    "regule (psycholog/inny)"
                )
            is_psychiatra = False
        else:
            is_psychiatra = employees_store.is_psychiatra(employee)

        if visit.status is None:
            log(
                f"Brak rozpoznanego statusu dla {visit.patient_name} "
                f"({visit.appointment_date.strftime('%d.%m.%Y')}) - wizyta pominieta"
            )
            continue

        if not _status_counts(visit.status, is_psychiatra):
            continue

        by_doctor.setdefault(visit.doctor or "Nieznany", []).append(visit)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    summary_ws = wb.create_sheet(title="Podsumowanie")
    summary_ws.append(["Lekarz", "Suma (zl)"])

    used_names: set[str] = {"Podsumowanie"}
    grand_total = 0.0

    for doctor in sorted(by_doctor):
        doctor_visits = sorted(by_doctor[doctor], key=lambda v: (v.appointment_date, v.appointment_time))
        ws = wb.create_sheet(title=_safe_sheet_name(doctor, used_names))
        ws.append(["Data", "Godzina", "Pacjent", "Cena (zl)", "Uwagi"])

        total = 0.0
        for visit in doctor_visits:
            ws.append([
                visit.appointment_date.strftime("%d.%m.%Y"),
                visit.appointment_time.strftime("%H:%M"),
                visit.patient_name,
                visit.price if visit.price is not None else None,
                visit.price_note or "",
            ])
            if visit.price is not None:
                total += visit.price

        ws.append(["", "", "Suma:", total, ""])
        summary_ws.append([doctor, total])
        grand_total += total

    summary_ws.append(["Razem", grand_total])
    wb.move_sheet("Podsumowanie", offset=-len(wb.sheetnames))
    wb.save(output_path)
