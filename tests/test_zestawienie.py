import openpyxl
import pytest

import employees_store
import zestawienie


def _write_report(tmp_path, content: str):
    path = tmp_path / "raport.csv"
    path.write_bytes(content.encode("utf-8-sig"))
    return path


def _fake_employees():
    return [
        {"name": "Doktor Psychiatra", "specialization": "Psychiatra"},
        {"name": "Doktor Inny", "specialization": "Psycholog"},
    ]


@pytest.fixture(autouse=True)
def _stub_employees(monkeypatch):
    monkeypatch.setattr(employees_store, "load", lambda *a, **k: _fake_employees())


def test_psychiatra_only_wykonane_counts(tmp_path):
    content = (
        "1 sierpnia 2026;Doktor Psychiatra\n"
        ";;;;;;;\n"
        "12345678901\n"
        "1;1;10:00;Jan Kowalski\n"
        "600111222\n"
        "100;Wykonane\n"
        "11111111111\n"
        "2;2;11:00;Anna Nowak\n"
        "600111223\n"
        "Nie zrealizowane\n"
    )
    input_path = _write_report(tmp_path, content)
    output_path = tmp_path / "out.xlsx"

    zestawienie.generate(input_path, output_path)

    wb = openpyxl.load_workbook(output_path)
    assert "Doktor Psychiatra" in wb.sheetnames
    sheet = wb["Doktor Psychiatra"]
    patients = [row[2] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert "Jan Kowalski" in patients
    assert "Anna Nowak" not in patients

    summary = dict((r[0], r[1]) for r in wb["Podsumowanie"].iter_rows(min_row=2, values_only=True))
    assert summary["Doktor Psychiatra"] == 100
    assert summary["Razem"] == 100


def test_other_counts_wykonane_and_nie_zrealizowane_but_not_rezygnacja(tmp_path):
    content = (
        "1 sierpnia 2026;Doktor Inny\n"
        ";;;;;;;\n"
        "12345678901\n"
        "1;1;10:00;Jan Kowalski\n"
        "600111222\n"
        "100;Wykonane\n"
        "11111111111\n"
        "2;2;11:00;Anna Nowak\n"
        "600111223\n"
        "Nie zrealizowane\n"
        "22222222222\n"
        "3;3;12:00;Piotr Zielinski\n"
        "600111224\n"
        "Rezygnacja z wykonania\n"
    )
    input_path = _write_report(tmp_path, content)
    output_path = tmp_path / "out.xlsx"

    zestawienie.generate(input_path, output_path)

    wb = openpyxl.load_workbook(output_path)
    sheet = wb["Doktor Inny"]
    patients = [row[2] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert "Jan Kowalski" in patients
    assert "Anna Nowak" in patients
    assert "Piotr Zielinski" not in patients


def test_unmatched_doctor_defaults_to_other_rule_and_warns_once(tmp_path):
    content = (
        "1 sierpnia 2026;Doktor Nieznany\n"
        ";;;;;;;\n"
        "12345678901\n"
        "1;1;10:00;Jan Kowalski\n"
        "600111222\n"
        "Nie zrealizowane\n"
        "11111111111\n"
        "2;2;11:00;Anna Nowak\n"
        "600111223\n"
        "Nie zrealizowane\n"
    )
    input_path = _write_report(tmp_path, content)
    output_path = tmp_path / "out.xlsx"

    warnings = []
    zestawienie.generate(input_path, output_path, log=warnings.append)

    unmatched_warnings = [w for w in warnings if "Nierozpoznany lekarz" in w]
    assert len(unmatched_warnings) == 1

    wb = openpyxl.load_workbook(output_path)
    patients = [row[2] for row in wb["Doktor Nieznany"].iter_rows(min_row=2, values_only=True)]
    assert "Jan Kowalski" in patients
    assert "Anna Nowak" in patients


def test_missing_status_column_is_rejected(tmp_path):
    # Old (reminders) template has no Status column at all.
    content = (
        "1 sierpnia 2026;Doktor Psychiatra\n"
        ";;;;;;;\n"
        "12345678901\n"
        "1;1;10:00;Jan Kowalski\n"
        "600111222\n"
        "150\n"
    )
    input_path = _write_report(tmp_path, content)
    output_path = tmp_path / "out.xlsx"

    with pytest.raises(ValueError):
        zestawienie.generate(input_path, output_path)
